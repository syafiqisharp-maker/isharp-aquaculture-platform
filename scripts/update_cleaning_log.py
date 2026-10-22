import os
import re
import sys
import glob
import zipfile
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# Paths
DOWNLOADS_DIR = r"C:\Users\syafiq\Downloads"
DOCS_FALLBACK = r"C:\Users\syafiq\Documents\WhatsApp Chat with Cleaning External  Inter\WhatsApp Chat with Cleaning External & Inter.txt"
TARGET_DIR = r"C:\Users\syafiq\OneDrive - BLUE ARCHIPELAGO BERHAD\Shared Documents - iSHARP Data Management Team\iSharp Data"
TARGET_FILE = os.path.join(TARGET_DIR, "Pond_Cleaning_Log.xlsx")

def find_latest_chat_file():
    """Finds the newest WhatsApp export zip or txt in Downloads, or falls back to Documents."""
    zip_files = glob.glob(os.path.join(DOWNLOADS_DIR, "*Cleaning*.zip")) + \
                glob.glob(os.path.join(DOWNLOADS_DIR, "*WhatsApp*.zip"))
    
    txt_files = glob.glob(os.path.join(DOWNLOADS_DIR, "*Cleaning*.txt")) + \
                glob.glob(os.path.join(DOWNLOADS_DIR, "*WhatsApp*.txt"))
    
    all_files = zip_files + txt_files
    if all_files:
        latest = max(all_files, key=os.path.getmtime)
        return latest
    
    if os.path.exists(DOCS_FALLBACK):
        return DOCS_FALLBACK
    
    return None

def extract_chat_text(file_path):
    """Extracts raw text from a .zip or .txt file."""
    if file_path.lower().endswith(".zip"):
        with zipfile.ZipFile(file_path, 'r') as z:
            txt_names = [n for n in z.namelist() if n.lower().endswith('.txt')]
            if not txt_names:
                raise ValueError(f"No .txt file found inside {file_path}")
            best_txt = max(txt_names, key=lambda n: z.getinfo(n).file_size)
            with z.open(best_txt) as f:
                return f.read().decode('utf-8', errors='ignore'), file_path
    else:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read(), file_path

def format_pond_name(pond_raw):
    """Formats pond names: 010107 -> 01.01.07, P091803 -> 09.18.03, SSC 6 -> SSC 6."""
    pond_clean = pond_raw.strip()
    digits = re.sub(r'\D', '', pond_clean)
    if len(digits) == 6:
        return f"{digits[:2]}.{digits[2:4]}.{digits[4:]}"
    return re.sub(r'\s+', ' ', pond_clean).upper()

def parse_date_obj(date_raw):
    """Parses date string into a real Python datetime.date object for true Excel sorting."""
    m = re.search(r'(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})', date_raw)
    if m:
        day = int(m.group(1))
        month = int(m.group(2))
        year = int(m.group(3))
        if year < 100:
            year += 2000
        try:
            return datetime.date(year, month, day)
        except ValueError:
            pass
    return None

def standardize_team_name(team_raw, pond_formatted):
    """
    Standardizes team names:
    - If user entered generic 'PO', 'Po', 'P.O', or 'P0':
      Derive the team number from the first two digits of the pond name (e.g. 07.14.01 -> PO 7).
    - If user entered 'PO 8' or 'PO8' -> 'PO 8'
    - Internal teams: 'int 1', 'INT 1' -> 'Int 1'
    - External teams: 'ext 2', 'EXT 2' -> 'Ext 2'
    """
    t_clean = re.sub(r'\s+', ' ', team_raw.strip())
    
    # Internal teams
    m = re.match(r'^int(?:ernal)?\s*(\d+)$', t_clean, re.I)
    if m:
        return f"Int {m.group(1)}"
    if re.match(r'^int(?:ernal)?$', t_clean, re.I):
        return "Int"
        
    # External teams
    m = re.match(r'^ext(?:ernal)?\s*(\d+)$', t_clean, re.I)
    if m:
        return f"Ext {m.group(1)}"
    if re.match(r'^ext(?:ernal)?$', t_clean, re.I):
        return "Ext"
        
    # Generic PO / Po / P.O / P0 -> derive module number from pond name
    if re.match(r'^p\.?[o0]\.?$', t_clean, re.I):
        # 1. Standard pond format XX.YY.ZZ (e.g., 07.14.01 -> 07 -> PO 7)
        m_pond = re.match(r'^(\d{2})\.', pond_formatted)
        if m_pond:
            return f"PO {int(m_pond.group(1))}"
            
        # 2. Starting with 2 digits (e.g., 0816TP -> 08 -> PO 8)
        m_start = re.match(r'^(\d{2})', pond_formatted)
        if m_start:
            return f"PO {int(m_start.group(1))}"
            
        # 3. SSC Module (e.g., SSC 6 -> PO 6, SSC M8 -> PO 8)
        m_ssc = re.search(r'SSC\s*M?0?(\d+)', pond_formatted, re.I)
        if m_ssc:
            return f"PO {int(m_ssc.group(1))}"
            
        return "PO"

    # PO teams with explicit number (e.g., PO1, PO 8, Po 3)
    m = re.match(r'^p\.?[o0]\.?\s*(\d+)$', t_clean, re.I)
    if m:
        return f"PO {m.group(1)}"
        
    # Pre / Post Harvest
    if re.search(r'pre\s*harvest\s*a|pre-?a', t_clean, re.I):
        return "Pre Harvest A"
    if re.search(r'pre\s*harvest\s*b|pre-?b', t_clean, re.I):
        return "Pre Harvest B"
        
    return t_clean

def parse_inspections(text):
    """Parses all inspection cards from chat text."""
    pattern = re.compile(
        r'Date:\s*([^\n\r]+)[\s\S]*?'
        r'Requestor:\s*([^\n\r]+)[\s\S]*?'
        r'Team:\s*([^\n\r]+)[\s\S]*?'
        r'▶️\s*([^-–\n]+)[-–]\s*([^\n\r]+)',
        re.IGNORECASE
    )
    
    records = []
    for date_raw, req_raw, team_raw, pond_raw, status_raw in pattern.findall(text):
        date_obj = parse_date_obj(date_raw)
        if not date_obj:
            continue
            
        pond_clean = format_pond_name(pond_raw)
        team_clean = standardize_team_name(team_raw, pond_clean)
        req_clean = req_raw.strip()
        
        is_passed = 'pass' in status_raw.lower()
        status_clean = "Passed" if is_passed else "Failed"
        
        records.append({
            "DateObj": date_obj,
            "Pond": pond_clean,
            "Team": team_clean,
            "Requestor": req_clean,
            "Status": status_clean
        })
    return records

def normalize_date_key(val):
    """Normalizes cell date value into (year, month, day) tuple for deduplication."""
    if isinstance(val, (datetime.date, datetime.datetime)):
        return (val.year, val.month, val.day)
    if isinstance(val, str):
        d = parse_date_obj(val)
        if d:
            return (d.year, d.month, d.day)
    return str(val)

def update_excel_log(records):
    """Updates or creates the Pond_Cleaning_Log.xlsx file with real dates and Excel Table."""
    headers = ["Date", "Pond", "Team", "Requestor", "Status"]
    os.makedirs(TARGET_DIR, exist_ok=True)
    
    existing_keys = set()
    wb = None
    ws = None
    
    if os.path.exists(TARGET_FILE):
        try:
            wb = openpyxl.load_workbook(TARGET_FILE)
            ws = wb.active
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and row[0] is not None and row[1] is not None:
                    d_key = normalize_date_key(row[0])
                    p_key = str(row[1]).strip()
                    s_key = str(row[4] if len(row) > 4 else '').strip()
                    existing_keys.add((d_key, p_key, s_key))
        except PermissionError:
            print("\n[ERROR] Pond_Cleaning_Log.xlsx is currently OPEN in Excel!")
            print("        Please CLOSE the file in Excel and run the update again.\n")
            return -1, -1
        except Exception as e:
            print(f"Warning reading existing Excel: {e}. Creating fresh workbook.")
            wb = None
            
    if wb is None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Cleaning Inspections"
        ws.append(headers)
    
    new_count = 0
    font_regular = Font(name="Calibri", size=11)
    font_passed = Font(name="Calibri", size=11, bold=True, color="1E4620")
    fill_passed = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    font_failed = Font(name="Calibri", size=11, bold=True, color="9C0006")
    fill_failed = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    border_thin = Side(border_style="thin", color="E0E0E0")
    row_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    for r in records:
        d_key = (r["DateObj"].year, r["DateObj"].month, r["DateObj"].day)
        key = (d_key, r["Pond"], r["Status"])
        if key not in existing_keys:
            existing_keys.add(key)
            row_idx = ws.max_row + 1
            
            date_cell = ws.cell(row=row_idx, column=1, value=r["DateObj"])
            date_cell.number_format = 'dd/mm/yyyy'
            date_cell.alignment = Alignment(horizontal="center")
            
            ws.cell(row=row_idx, column=2, value=r["Pond"]).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=3, value=r["Team"]).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=4, value=r["Requestor"]).alignment = Alignment(horizontal="left")
            
            status_cell = ws.cell(row=row_idx, column=5, value=r["Status"])
            status_cell.alignment = Alignment(horizontal="center")
            if r["Status"] == "Passed":
                status_cell.font = font_passed
                status_cell.fill = fill_passed
            else:
                status_cell.font = font_failed
                status_cell.fill = fill_failed
                
            for col in range(1, 6):
                c = ws.cell(row=row_idx, column=col)
                if col != 5:
                    c.font = font_regular
                c.border = row_border
            ws.row_dimensions[row_idx].height = 20
            new_count += 1

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for col in range(1, 6):
        c = ws.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    max_row = max(ws.max_row, 2)
    table_range = f"A1:E{max_row}"
    
    if ws.tables:
        for t_name in list(ws.tables.keys()):
            ws.tables[t_name].ref = table_range
    else:
        tab = Table(displayName="PondCleaningTable", ref=table_range)
        style = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )
        tab.tableStyleInfo = style
        ws.add_table(tab)
            
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)
        
    try:
        wb.save(TARGET_FILE)
    except PermissionError:
        print("\n[ERROR] Pond_Cleaning_Log.xlsx is currently OPEN in Excel!")
        print("        Please CLOSE the file in Excel and run the update again.\n")
        return -1, -1

    return new_count, len(existing_keys)

def rebuild_full_log():
    """Rebuilds the entire log from scratch using all historical records to fix formats."""
    print("[REBUILD] Rebuilding log from scratch with proper Date data types and PO Module Team numbers...")
    if os.path.exists(TARGET_FILE):
        try:
            os.remove(TARGET_FILE)
        except PermissionError:
            print("[ERROR] Cannot delete old file because it is currently OPEN in Excel!")
            print("        Please close it and try again.")
            return False

    text, _ = extract_chat_text(DOCS_FALLBACK)
    records = parse_inspections(text)
    # Sort chronologically by date
    records.sort(key=lambda r: (r["DateObj"], r["Pond"]))
    
    new_added, total = update_excel_log(records)
    print(f"[REBUILD] Successfully imported {new_added} records with PO team numbers mapped from pond prefix!")
    return True

def main():
    sys.stdout.reconfigure(encoding='utf-8')
    print("=" * 60)
    print("     POND CLEANING LOG - 1-CLICK UPDATE TOOL")
    print("=" * 60)
    
    if len(sys.argv) > 1 and sys.argv[1] == "--rebuild":
        rebuild_full_log()
        return 0

    source_file = find_latest_chat_file()
    if not source_file:
        print("[ERROR] No WhatsApp export file (.zip or .txt) found.")
        return 1
        
    print(f"[FOUND] Source export: {os.path.basename(source_file)}")
    print(f"        Path: {source_file}")
    
    try:
        text, actual_file = extract_chat_text(source_file)
    except Exception as e:
        print(f"[ERROR] Failed to read source file: {e}")
        return 1
        
    records = parse_inspections(text)
    print(f"[PARSED] Found {len(records)} inspection records in export.")
    
    new_added, total_entries = update_excel_log(records)
    if new_added == -1:
        return 1
        
    print(f"[SUCCESS] Added {new_added} new records to Excel.")
    print(f"[STATUS] Total unique records in log: {total_entries}")
    print(f"[TABLE] Official Excel Table created with sort & filter arrows.")
    print(f"[SAVED] {TARGET_FILE}")
    print("=" * 60)
    return 0

if __name__ == "__main__":
    exit_code = main()
    sys.exit(exit_code)
