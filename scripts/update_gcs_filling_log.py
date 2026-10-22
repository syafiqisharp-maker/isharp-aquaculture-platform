import os
import re
import sys
import glob
import zipfile
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

DOWNLOADS_DIR = r"C:\Users\syafiq\Downloads"
TARGET_DIR = r"C:\Users\syafiq\OneDrive - BLUE ARCHIPELAGO BERHAD\Shared Documents - iSHARP Data Management Team\iSharp Data"
TARGET_FILE = os.path.join(TARGET_DIR, "Pond_Water_Filling_Log.xlsx")

def find_latest_gcs_file():
    """Finds the newest GCS export zip or txt in Downloads."""
    zip_files = glob.glob(os.path.join(DOWNLOADS_DIR, "*GCS*.zip"))
    txt_files = glob.glob(os.path.join(DOWNLOADS_DIR, "*GCS*.txt"))
    all_files = zip_files + txt_files
    if all_files:
        return max(all_files, key=os.path.getmtime)
    return None

def extract_chat_text(file_path):
    """Extracts raw text from a .zip or .txt file."""
    if file_path.lower().endswith(".zip"):
        with zipfile.ZipFile(file_path, 'r') as z:
            txt_names = [n for n in z.namelist() if n.lower().endswith('.txt')]
            if not txt_names:
                raise ValueError(f"No .txt file found inside {file_path}")
            best_txt = max(txt_names, key=lambda n: z.getinfo(n).file_size)
            with z.open(best_txt) as f:
                return f.read().decode('utf-8', errors='ignore'), file_path
    else:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read(), file_path

def parse_wa_dt(dt_str):
    """Parses WhatsApp timestamp '[M/D/YY, H:MM:SS AM/PM]' into datetime object."""
    m = re.search(r'(\d{1,2})[/-](\d{1,2})[/-](\d{2,4}),?\s*(\d{1,2}):(\d{2}):(\d{2})\s*(AM|PM|am|pm)?', dt_str)
    if not m:
        return None
    month, day, year = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if year < 100:
        year += 2000
    hour, minute, sec = int(m.group(4)), int(m.group(5)), int(m.group(6))
    ampm = m.group(7)
    if ampm:
        ampm = ampm.upper()
        if ampm == 'PM' and hour < 12:
            hour += 12
        elif ampm == 'AM' and hour == 12:
            hour = 0
    return datetime.datetime(year, month, day, hour, minute, sec)

def parse_gcs_reports(text):
    """Extracts and sorts all chronological GCS Daily Status reports."""
    pattern = re.compile(
        r'^\[(\d{1,2}/\d{1,2}/\d{2,4},\s*\d{1,2}:\d{2}:\d{2}\s*(?:AM|PM|am|pm)?)\]\s*([^:\n]+):',
        re.MULTILINE
    )
    matches = list(pattern.finditer(text))
    reports = []
    
    for i in range(len(matches)):
        start = matches[i].start()
        end = matches[i+1].start() if i+1 < len(matches) else len(text)
        header_time = matches[i].group(1)
        dt = parse_wa_dt(header_time)
        body = text[matches[i].end():end].strip()
        if ('status pumping' in body.lower() or 'report harian status gcs' in body.lower()) and dt:
            reports.append((dt, body))
            
    reports.sort(key=lambda x: x[0])
    return reports

def parse_report_ponds(body):
    """Parses ponds and their status symbols under each GCS section."""
    current_gcs = "GCS 1"
    ponds = []
    for line in body.splitlines():
        line_clean = line.strip()
        if re.search(r'gcs\s*2', line_clean, re.I):
            current_gcs = "GCS 2"
        elif re.search(r'gcs\s*1', line_clean, re.I):
            current_gcs = "GCS 1"
            
        if any(k in line_clean.lower() for k in ['gas', 'dosage', 'pump', 'booster', 'tong', 'empty', 'tarikh', 'report', 'catatan', 'used']):
            continue
            
        sym = None
        if '🔛' in line_clean:
            sym = '🔛'
        elif '✅' in line_clean:
            sym = '✅'
            
        # 6-digit ponds
        m6 = re.search(r'\b(\d{6})\b', line_clean)
        if m6:
            code = m6.group(1)
            formatted = f"{code[:2]}.{code[2:4]}.{code[4:]}"
            ponds.append((current_gcs, formatted, sym))
            continue
            
        # Reservoirs (TP / TA)
        mtp = re.search(r'\b(T[PA]\s*\d+)\b', line_clean, re.I)
        if mtp:
            clean_tp = re.sub(r'\s+', ' ', mtp.group(1).upper())
            ponds.append((current_gcs, clean_tp, sym))
            continue
            
    return ponds

def extract_filling_cycles(reports):
    """
    Runs state machine across reports to capture complete filling cycles
    and ponds currently still in progress.
    """
    completed_cycles = []
    active_fillings = {}
    recently_completed = {}

    for dt, body in reports:
        p_list = parse_report_ponds(body)
        ponds_in_report = set()
        
        for gcs, pond, sym in p_list:
            ponds_in_report.add(pond)
            
            if sym == '🔛':
                if pond in recently_completed:
                    del recently_completed[pond]
                if pond not in active_fillings:
                    active_fillings[pond] = {
                        'gcs': gcs,
                        'start_dt': dt,
                        'last_seen_dt': dt
                    }
                else:
                    active_fillings[pond]['last_seen_dt'] = dt
                    active_fillings[pond]['gcs'] = gcs
                    
            elif sym == '✅':
                if pond in recently_completed:
                    # Ignore duplicate report of already completed pond
                    continue
                    
                if pond in active_fillings:
                    start_dt = active_fillings[pond]['start_dt']
                    dur_days = round((dt - start_dt).total_seconds() / 86400.0, 2)
                    completed_cycles.append({
                        'pond': pond,
                        'gcs': active_fillings[pond]['gcs'],
                        'start_dt': start_dt,
                        'full_dt': dt,
                        'duration_days': dur_days,
                        'status': 'Completed ✅',
                        'note': 'Confirmed ✅'
                    })
                    del active_fillings[pond]
                    recently_completed[pond] = dt
                else:
                    # Started prior to export range
                    completed_cycles.append({
                        'pond': pond,
                        'gcs': gcs,
                        'start_dt': None,
                        'full_dt': dt,
                        'duration_days': None,
                        'status': 'Completed ✅',
                        'note': 'Confirmed ✅ (Existing at Start)'
                    })
                    recently_completed[pond] = dt

        # Detect operator forgetting ✅ and removing pond
        to_delete = []
        for pond, data in active_fillings.items():
            if pond not in ponds_in_report:
                start_dt = data['start_dt']
                dur_days = round((dt - start_dt).total_seconds() / 86400.0, 2)
                completed_cycles.append({
                    'pond': pond,
                    'gcs': data['gcs'],
                    'start_dt': start_dt,
                    'full_dt': dt,
                    'duration_days': dur_days,
                    'status': 'Completed ✅',
                    'note': 'Implied (Removed without ✅)'
                })
                to_delete.append(pond)
                recently_completed[pond] = dt
        for p in to_delete:
            del active_fillings[p]

    # Any remaining active fillings are currently still in progress
    in_progress_records = []
    for pond, data in active_fillings.items():
        in_progress_records.append({
            'pond': pond,
            'gcs': data['gcs'],
            'start_dt': data['start_dt'],
            'full_dt': None,
            'duration_days': None,
            'status': 'In Progress 🔛',
            'note': 'Still Filling'
        })
        
    return completed_cycles, in_progress_records

def update_excel_filling_log(completed_cycles, in_progress_records):
    """Writes filling records into Pond_Water_Filling_Log.xlsx."""
    headers = ["Pond", "GCS Unit", "Start Date & Time", "Full Date & Time", "Duration (Days)", "Status", "Completion Note"]
    os.makedirs(TARGET_DIR, exist_ok=True)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Water Filling Log"
    ws.append(headers)
    
    # Headers style
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for col in range(1, 8):
        c = ws.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25
    
    font_regular = Font(name="Calibri", size=11)
    font_passed = Font(name="Calibri", size=11, bold=True, color="1E4620")
    fill_passed = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    font_prog = Font(name="Calibri", size=11, bold=True, color="8A6D3B")
    fill_prog = PatternFill(start_color="FCF8E3", end_color="FCF8E3", fill_type="solid")
    border_thin = Side(border_style="thin", color="E0E0E0")
    row_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    all_records = completed_cycles + in_progress_records
    # Sort chronologically by start_dt if present, otherwise full_dt
    all_records.sort(key=lambda r: (r['start_dt'] or r['full_dt'] or datetime.datetime.min, r['pond']))
    
    for r in all_records:
        row_idx = ws.max_row + 1
        ws.cell(row=row_idx, column=1, value=r['pond']).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=2, value=r['gcs']).alignment = Alignment(horizontal="center")
        
        # Start datetime
        c_start = ws.cell(row=row_idx, column=3, value=r['start_dt'])
        if r['start_dt']:
            c_start.number_format = 'dd/mm/yyyy hh:mm'
        c_start.alignment = Alignment(horizontal="center")
        
        # Full datetime
        c_full = ws.cell(row=row_idx, column=4, value=r['full_dt'])
        if r['full_dt']:
            c_full.number_format = 'dd/mm/yyyy hh:mm'
        c_full.alignment = Alignment(horizontal="center")
        
        # Duration days (numeric for formula/sorting)
        c_dur = ws.cell(row=row_idx, column=5, value=r['duration_days'])
        if r['duration_days'] is not None:
            c_dur.number_format = '0.00'
        c_dur.alignment = Alignment(horizontal="right")
        
        # Status
        c_stat = ws.cell(row=row_idx, column=6, value=r['status'])
        c_stat.alignment = Alignment(horizontal="center")
        if 'Completed' in r['status']:
            c_stat.font = font_passed
            c_stat.fill = fill_passed
        else:
            c_stat.font = font_prog
            c_stat.fill = fill_prog
            
        # Note
        c_note = ws.cell(row=row_idx, column=7, value=r['note'])
        c_note.alignment = Alignment(horizontal="left")
        
        for col in range(1, 8):
            cell = ws.cell(row=row_idx, column=col)
            if col != 6:
                cell.font = font_regular
            cell.border = row_border
        ws.row_dimensions[row_idx].height = 20

    max_row = max(ws.max_row, 2)
    table_range = f"A1:G{max_row}"
    tab = Table(displayName="PondWaterFillingTable", ref=table_range)
    style = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )
    tab.tableStyleInfo = style
    ws.add_table(tab)
    
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)
        
    try:
        wb.save(TARGET_FILE)
    except PermissionError:
        print("\n[ERROR] Pond_Water_Filling_Log.xlsx is currently OPEN in Excel!")
        print("        Please CLOSE the file in Excel and run the update again.\n")
        return -1
        
    return len(completed_cycles), len(in_progress_records)

def main():
    sys.stdout.reconfigure(encoding='utf-8')
    print("=" * 60)
    print("     POND WATER FILLING LOG - GCS UPDATE TOOL")
    print("=" * 60)
    
    source_file = find_latest_gcs_file()
    if not source_file:
        print(f"[SKIP] No GCS export file found in: {DOWNLOADS_DIR}")
        return 0
        
    print(f"[FOUND] GCS source export: {os.path.basename(source_file)}")
    print(f"        Path: {source_file}")
    
    try:
        text, actual_file = extract_chat_text(source_file)
    except Exception as e:
        print(f"[ERROR] Failed to read GCS file: {e}")
        return 1
        
    reports = parse_gcs_reports(text)
    print(f"[PARSED] Found {len(reports)} GCS status reports in export.")
    
    completed, in_prog = extract_filling_cycles(reports)
    res = update_excel_filling_log(completed, in_prog)
    if res == -1:
        return 1
        
    print(f"[SUCCESS] Recorded {completed[0] if isinstance(completed, tuple) else res[0]} completed filling cycles.")
    print(f"[ACTIVE]  Currently filling right now: {res[1]} ponds.")
    print(f"[SAVED]   {TARGET_FILE}")
    print("=" * 60)
    return 0

if __name__ == "__main__":
    exit_code = main()
    sys.exit(exit_code)
