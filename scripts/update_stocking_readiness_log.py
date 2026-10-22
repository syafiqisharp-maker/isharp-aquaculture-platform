import os
import re
import sys
import glob
import zipfile
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

DOWNLOADS_DIR = r"C:\Users\syafiq\Downloads"
TARGET_DIR = r"C:\Users\syafiq\OneDrive - BLUE ARCHIPELAGO BERHAD\Shared Documents - iSHARP Data Management Team\iSharp Data"
TARGET_FILE = os.path.join(TARGET_DIR, "Pond_Stocking_Readiness_Log.xlsx")

def find_latest_production_file():
    """Finds the newest Production export zip or txt in Downloads."""
    patterns = ["*Production*.zip", "*Stocking*.zip", "*Production*.txt", "*Stocking*.txt"]
    all_files = []
    for pat in patterns:
        all_files.extend(glob.glob(os.path.join(DOWNLOADS_DIR, pat)))
    if all_files:
        return max(all_files, key=os.path.getmtime)
    return None

def extract_chat_text(file_path):
    """Extracts raw text from a .zip or .txt file."""
    if file_path.lower().endswith(".zip"):
        with zipfile.ZipFile(file_path, 'r') as z:
            txt_names = [n for n in z.namelist() if n.lower().endswith('.txt')]
            if not txt_names:
                raise ValueError(f"No .txt file found inside {file_path}")
            best_txt = max(txt_names, key=lambda n: z.getinfo(n).file_size)
            with z.open(best_txt) as f:
                return f.read().decode('utf-8', errors='ignore'), file_path
    else:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read(), file_path

def parse_date_obj(date_raw):
    """Parses date string into a real Python datetime.date object for true Excel sorting."""
    m = re.search(r'(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})', date_raw)
    if m:
        day = int(m.group(1))
        month = int(m.group(2))
        year = int(m.group(3))
        if year < 100:
            year += 2000
        try:
            return datetime.date(year, month, day)
        except ValueError:
            pass
    return None

def parse_stocking_records(text):
    """Parses all pond stocking preparation inspections from chat text."""
    pattern = re.compile(
        r'^\[(\d{1,2}/\d{1,2}/\d{2,4},\s*\d{1,2}:\d{2}:\d{2}\s*(?:AM|PM|am|pm)?)\]\s*([^:\n]+):',
        re.MULTILINE
    )
    matches = list(pattern.finditer(text))
    
    inspection_posts = []
    for i in range(len(matches)):
        start = matches[i].start()
        end = matches[i+1].start() if i+1 < len(matches) else len(text)
        header_time = matches[i].group(1)
        sender = matches[i].group(2)
        body = text[matches[i].end():end].strip()
        if 'stocking preparation' in body.lower() or 'grow out book' in body.lower():
            inspection_posts.append((header_time, sender, body))
            
    records = []
    for ht, sender, body in inspection_posts:
        m_date = re.search(r'Date:\s*([^\n\r]+)', body, re.I)
        date_obj = parse_date_obj(m_date.group(1)) if m_date else parse_date_obj(ht)
        if not date_obj:
            continue
            
        m_req = re.search(r'Requestor:\s*([^\n\r]+)', body, re.I)
        req = m_req.group(1).strip() if m_req else sender.strip()
        
        pond_blocks = re.split(r'▶️', body)[1:]
        for block in pond_blocks:
            first_line = block.splitlines()[0].strip()
            m_pond = re.search(r'(\d{6})', first_line)
            if not m_pond:
                continue
                
            raw_code = m_pond.group(1)
            formatted_pond = f"{raw_code[:2]}.{raw_code[2:4]}.{raw_code[4:]}"
            
            # Determine readiness status & pending reason
            status = "Ready for Stocking ✅"
            pending_reason = "-"
            
            first_line_clean = first_line.replace(raw_code, '')
            if '❌' in first_line or '⚠️' in first_line or 'pending' in first_line.lower():
                status = "Pending ❌"
                cleaned = re.sub(r'^(?:[▶️⚠️❌\s]|\d{6}|\d{2}\.\d{2}\.\d{2})+', '', first_line).strip()
                cleaned = re.sub(r'^[❌⚠️\s]+', '', cleaned).strip()
                pending_reason = cleaned if cleaned else "Pending inspection"
            elif '❌' in block:
                status = "Pending ❌"
                # Look for line with ❌ in checklist
                for bline in block.splitlines():
                    if '❌' in bline:
                        pending_reason = bline.replace('❌', '').strip()
                        break
                        
            # DOP
            m_dop = re.search(r'DOP\s*:\s*([^\n\r]+)', block, re.I)
            dop = m_dop.group(1).strip() if m_dop else "-"
            
            records.append({
                "Date": date_obj,
                "Pond": formatted_pond,
                "Requestor": req,
                "Status": status,
                "PendingReason": pending_reason,
                "DOP": dop
            })
            
    return records

def normalize_date_key(val):
    """Normalizes cell date value into (year, month, day) tuple for deduplication."""
    if isinstance(val, (datetime.date, datetime.datetime)):
        return (val.year, val.month, val.day)
    if isinstance(val, str):
        d = parse_date_obj(val)
        if d:
            return (d.year, d.month, d.day)
    return str(val)

def update_excel_stocking_log(records):
    """Updates or creates the Pond_Stocking_Readiness_Log.xlsx file."""
    headers = ["Date", "Pond", "Requestor", "Readiness Status", "Pending Reason", "DOP"]
    os.makedirs(TARGET_DIR, exist_ok=True)
    
    existing_keys = set()
    wb = None
    ws = None
    
    if os.path.exists(TARGET_FILE):
        try:
            wb = openpyxl.load_workbook(TARGET_FILE)
            ws = wb.active
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and row[0] is not None and row[1] is not None:
                    d_key = normalize_date_key(row[0])
                    p_key = str(row[1]).strip()
                    s_key = str(row[3] if len(row) > 3 else '').strip()
                    existing_keys.add((d_key, p_key, s_key))
        except PermissionError:
            print("\n[ERROR] Pond_Stocking_Readiness_Log.xlsx is currently OPEN in Excel!")
            print("        Please CLOSE the file in Excel and run the update again.\n")
            return -1, -1
        except Exception as e:
            print(f"Warning reading existing Excel: {e}. Creating fresh workbook.")
            wb = None
            
    if wb is None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Stocking Readiness"
        ws.append(headers)
        
    new_count = 0
    font_regular = Font(name="Calibri", size=11)
    font_ready = Font(name="Calibri", size=11, bold=True, color="1E4620")
    fill_ready = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    font_pending = Font(name="Calibri", size=11, bold=True, color="9C0006")
    fill_pending = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    border_thin = Side(border_style="thin", color="E0E0E0")
    row_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    for r in records:
        d_key = (r["Date"].year, r["Date"].month, r["Date"].day)
        key = (d_key, r["Pond"], r["Status"])
        if key not in existing_keys:
            existing_keys.add(key)
            row_idx = ws.max_row + 1
            
            # Real date
            c_date = ws.cell(row=row_idx, column=1, value=r["Date"])
            c_date.number_format = 'dd/mm/yyyy'
            c_date.alignment = Alignment(horizontal="center")
            
            ws.cell(row=row_idx, column=2, value=r["Pond"]).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=3, value=r["Requestor"]).alignment = Alignment(horizontal="center")
            
            # Status
            c_stat = ws.cell(row=row_idx, column=4, value=r["Status"])
            c_stat.alignment = Alignment(horizontal="center")
            if "Ready" in r["Status"]:
                c_stat.font = font_ready
                c_stat.fill = fill_ready
            else:
                c_stat.font = font_pending
                c_stat.fill = fill_pending
                
            # Reason & DOP
            ws.cell(row=row_idx, column=5, value=r["PendingReason"]).alignment = Alignment(horizontal="left")
            ws.cell(row=row_idx, column=6, value=r["DOP"]).alignment = Alignment(horizontal="center")
            
            for col in range(1, 7):
                cell = ws.cell(row=row_idx, column=col)
                if col != 4:
                    cell.font = font_regular
                cell.border = row_border
            ws.row_dimensions[row_idx].height = 20
            new_count += 1
            
    # Format headers
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for col in range(1, 7):
        c = ws.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    # Table
    max_row = max(ws.max_row, 2)
    table_range = f"A1:F{max_row}"
    if ws.tables:
        for t_name in list(ws.tables.keys()):
            ws.tables[t_name].ref = table_range
    else:
        tab = Table(displayName="PondStockingTable", ref=table_range)
        style = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )
        tab.tableStyleInfo = style
        ws.add_table(tab)
        
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)
        
    try:
        wb.save(TARGET_FILE)
    except PermissionError:
        print("\n[ERROR] Pond_Stocking_Readiness_Log.xlsx is currently OPEN in Excel!")
        print("        Please CLOSE the file in Excel and run the update again.\n")
        return -1, -1

    return new_count, len(existing_keys)

def rebuild_full_log():
    """Rebuilds the log from scratch."""
    if os.path.exists(TARGET_FILE):
        try:
            os.remove(TARGET_FILE)
        except PermissionError:
            print("[ERROR] Cannot delete old file because it is OPEN in Excel!")
            return False
            
    source_file = find_latest_production_file()
    if not source_file:
        print("No source file found.")
        return False
        
    text, _ = extract_chat_text(source_file)
    records = parse_stocking_records(text)
    records.sort(key=lambda r: (r["Date"], r["Pond"]))
    new_added, total = update_excel_stocking_log(records)
    print(f"[REBUILD] Imported {new_added} records into {TARGET_FILE}")
    return True

def main():
    sys.stdout.reconfigure(encoding='utf-8')
    print("=" * 60)
    print("     POND STOCKING READINESS LOG - UPDATE TOOL")
    print("=" * 60)
    
    if len(sys.argv) > 1 and sys.argv[1] == "--rebuild":
        rebuild_full_log()
        return 0

    source_file = find_latest_production_file()
    if not source_file:
        print(f"[SKIP] No Production export file found in: {DOWNLOADS_DIR}")
        return 0
        
    print(f"[FOUND] Production source export: {os.path.basename(source_file)}")
    print(f"        Path: {source_file}")
    
    try:
        text, actual_file = extract_chat_text(source_file)
    except Exception as e:
        print(f"[ERROR] Failed to read Production file: {e}")
        return 1
        
    records = parse_stocking_records(text)
    records.sort(key=lambda r: (r["Date"], r["Pond"]))
    print(f"[PARSED] Found {len(records)} pond stocking preparation records.")
    
    new_added, total_entries = update_excel_stocking_log(records)
    if new_added == -1:
        return 1
        
    print(f"[SUCCESS] Added {new_added} new records to Excel.")
    print(f"[STATUS]  Total unique records in log: {total_entries}")
    print(f"[TABLE]   Official Excel Table created with sort & filter arrows.")
    print(f"[SAVED]   {TARGET_FILE}")
    print("=" * 60)
    return 0

if __name__ == "__main__":
    exit_code = main()
    sys.exit(exit_code)
