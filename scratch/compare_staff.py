import openpyxl
import json
import urllib.request

# 1. Fetch Supabase pond_staff
url = 'https://keappoukeagyzpoxkrru.supabase.co/rest/v1/pond_staff?select=*&limit=1000'
headers = {
    'apikey': 'sb_publishable_kObmQ9Ha4NLrl9vQXy5k5w_Ie9-EQ4s',
    'Authorization': 'Bearer sb_publishable_kObmQ9Ha4NLrl9vQXy5k5w_Ie9-EQ4s'
}
req = urllib.request.Request(url, headers=headers)
with urllib.request.urlopen(req) as resp:
    db_data = json.loads(resp.read().decode('utf-8'))

# 2. Read Excel
excel_path = r'c:\Users\syafiq\My Drive\Syafiq Water Quality Station Project\Pond Operations Management System\20261011 Current Pond worker.xlsx'
wb = openpyxl.load_workbook(excel_path)
sheet = wb.active

excel_records = []
for r in range(3, sheet.max_row + 1):
    s_no = sheet.cell(r, 2).value
    s_name = sheet.cell(r, 3).value
    s_pos = sheet.cell(r, 4).value
    if s_no is not None:
        excel_records.append({
            'staff_no': str(s_no).strip(),
            'staff_name': str(s_name).strip() if s_name else '',
            'staff_position': str(s_pos).strip() if s_pos else ''
        })

print(f'DB total count: {len(db_data)}')
print(f'Excel total count: {len(excel_records)}')

db_map = {str(d['staff_no']).strip(): d for d in db_data}
excel_map = {str(e['staff_no']).strip(): e for e in excel_records}

# In Excel but not in DB (New staff)
in_excel_not_db = [e for s_no, e in excel_map.items() if s_no not in db_map]

# In DB but not in Excel (Removed / missing staff)
in_db_not_excel = [d for s_no, d in db_map.items() if s_no not in excel_map]

# In both but attributes differ
differences = []
for s_no, e in excel_map.items():
    if s_no in db_map:
        d = db_map[s_no]
        diffs = {}
        if e['staff_name'] != d['staff_name']:
            diffs['name'] = {'excel': e['staff_name'], 'db': d['staff_name']}
        if e['staff_position'] != d['staff_position']:
            diffs['position'] = {'excel': e['staff_position'], 'db': d['staff_position']}
        if diffs:
            differences.append({'staff_no': s_no, 'diffs': diffs, 'excel': e, 'db': d})

print('\n==================================================')
print(f'1. IN EXCEL BUT NOT IN DB (NEW FROM HR) [{len(in_excel_not_db)}]')
print('==================================================')
for x in in_excel_not_db:
    print(f"Staff No: {x['staff_no']} | Name: {x['staff_name']} | Position: {x['staff_position']}")

print('\n==================================================')
print(f'2. IN DB BUT NOT IN EXCEL (NOT IN HR LIST) [{len(in_db_not_excel)}]')
print('==================================================')
for x in in_db_not_excel:
    print(f"Staff No: {x['staff_no']} | Name: {x['staff_name']} | Position: {x['staff_position']} | Active: {x.get('is_active')}")

print('\n==================================================')
print(f'3. IN BOTH BUT WITH DIFFERENCES [{len(differences)}]')
print('==================================================')
for x in differences:
    print(f"Staff No: {x['staff_no']}")
    for k, v in x['diffs'].items():
        print(f"   - {k.upper()}: DB = '{v['db']}'  <-->  Excel = '{v['excel']}'")
