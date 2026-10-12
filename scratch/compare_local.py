import openpyxl
import json

with open(r'c:\Users\syafiq\My Drive\Syafiq Water Quality Station Project\scratch\supabase_pond_staff.json', 'r', encoding='utf-8-sig') as f:
    db_data = json.load(f)

excel_path = r'c:\Users\syafiq\My Drive\Syafiq Water Quality Station Project\Pond Operations Management System\20261011 Current Pond worker.xlsx'
wb = openpyxl.load_workbook(excel_path)
sheet = wb.active

excel_records = []
for r in range(3, sheet.max_row + 1):
    s_no = sheet.cell(r, 2).value
    s_name = sheet.cell(r, 3).value
    s_pos = sheet.cell(r, 4).value
    if s_no is not None:
        excel_records.append({
            'staff_no': str(s_no).strip(),
            'staff_name': str(s_name).strip() if s_name else '',
            'staff_position': str(s_pos).strip() if s_pos else ''
        })

print(f'DB total count: {len(db_data)}')
print(f'Excel total count: {len(excel_records)}')

db_map = {str(d['staff_no']).strip(): d for d in db_data}
excel_map = {str(e['staff_no']).strip(): e for e in excel_records}

# 1. In Excel but not in DB
in_excel_not_db = [e for s_no, e in excel_map.items() if s_no not in db_map]

# 2. In DB but not in Excel
in_db_not_excel = [d for s_no, d in db_map.items() if s_no not in excel_map]

# 3. In both but different
differences = []
for s_no, e in excel_map.items():
    if s_no in db_map:
        d = db_map[s_no]
        diffs = {}
        # normalize whitespace
        e_name = ' '.join(e['staff_name'].split())
        d_name = ' '.join(d['staff_name'].split())
        e_pos = ' '.join(e['staff_position'].split())
        d_pos = ' '.join(d['staff_position'].split())
        if e_name.lower() != d_name.lower():
            diffs['name'] = {'excel': e['staff_name'], 'db': d['staff_name']}
        elif e_name != d_name:
            diffs['name_case_or_spacing'] = {'excel': e['staff_name'], 'db': d['staff_name']}
        if e_pos != d_pos:
            diffs['position'] = {'excel': e['staff_position'], 'db': d['staff_position']}
        if diffs:
            differences.append({'staff_no': s_no, 'diffs': diffs, 'excel': e, 'db': d})

result = {
    'summary': {
        'db_total': len(db_data),
        'excel_total': len(excel_records),
        'new_in_excel_count': len(in_excel_not_db),
        'missing_from_excel_count': len(in_db_not_excel),
        'differences_count': len(differences)
    },
    'new_in_excel': in_excel_not_db,
    'missing_from_excel': in_db_not_excel,
    'differences': differences
}

with open(r'c:\Users\syafiq\My Drive\Syafiq Water Quality Station Project\scratch\comparison_result.json', 'w', encoding='utf-8') as f:
    json.dump(result, f, indent=2, ensure_ascii=False)

print('Successfully written comparison_result.json')
