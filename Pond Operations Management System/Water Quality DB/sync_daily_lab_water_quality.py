#!/usr/bin/env python3
"""
iSHARP Aquaculture Platform — Daily Automated Laboratory Water Quality Sync
=============================================================================
Location:
  C:\\Users\\syafiq\\My Drive\\Syafiq Water Quality Station Project\\Pond Operations Management System\\Water Quality DB\\sync_daily_lab_water_quality.py

Target Database:
  Y:\\9. Database\\DATABASE COMBINE MONITORING 2024 - 2026 new.xlsx (Office Server)
  Fallback: Local Water Quality DB directory if server drive is temporarily disconnected.

Features:
  1. Safe Temp-Copy: Copies file to temp directory to avoid lock conflicts with lab personnel editing in Excel.
  2. Smart Incremental Sync: Looks back at recent records (last 30 days) to update and insert fresh lab readings in 2-5 seconds.
  3. Full Resync Flag: Supports '--full' command line argument to resync the entire 2024-2026 archive if needed.
  4. Automatic Cycle Matching: Accurately maps pond name + sampling date to its active/historical cycle pond_index.
  5. Logging: Writes timestamps and sync results to 'sync_lab_water_quality.log' in the same folder.
"""

import os
import sys

# Ensure immediate unbuffered console output
sys.stdout.reconfigure(line_buffering=True)

import json
import time
import shutil
import tempfile
import datetime
import urllib.request
import urllib.error
import openpyxl

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
LOG_FILE = os.path.join(SCRIPT_DIR, "sync_lab_water_quality.log")

# Primary source: Office Server Drive Y:
PRIMARY_EXCEL_PATH = r"Y:\9. Database\DATABASE COMBINE MONITORING 2024 - 2026 new.xlsx"
# Fallback source: Local copy
FALLBACK_EXCEL_PATH = os.path.join(SCRIPT_DIR, "DATABASE COMBINE MONITORING 2024 - 2026 new.xlsx")

SUPABASE_URL = "https://keappoukeagyzpoxkrru.supabase.co"
SUPABASE_KEY = "sb_publishable_kObmQ9Ha4NLrl9vQXy5k5w_Ie9-EQ4s"
BATCH_SIZE = 500

HEADERS = {
    "apikey": SUPABASE_KEY,
    "Authorization": f"Bearer {SUPABASE_KEY}",
    "Content-Type": "application/json",
    "Prefer": "resolution=merge-duplicates"
}

def log(msg):
    timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    formatted = f"[{timestamp}] {msg}"
    print(formatted)
    try:
        with open(LOG_FILE, "a", encoding="utf-8") as f:
            f.write(formatted + "\n")
    except Exception:
        pass

def safe_float(val):
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    val_str = str(val).strip().replace(",", "")
    if not val_str or val_str in ("-", "ND", "N/A", "nil", "NA", "None"):
        return None
    try:
        return float(val_str)
    except ValueError:
        clean = "".join(c for c in val_str if c.isdigit() or c == ".")
        try:
            return float(clean) if clean else None
        except ValueError:
            return None

def safe_int(val):
    flt = safe_float(val)
    return int(round(flt)) if flt is not None else None

def safe_date_str(val):
    if val is None:
        return None
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.strftime("%Y-%m-%d")
    val_str = str(val).strip()[:10]
    try:
        dt = datetime.datetime.strptime(val_str, "%Y-%m-%d")
        return dt.strftime("%Y-%m-%d")
    except ValueError:
        return None

def fetch_master_cycles(max_retries=5):
    log(f"Fetching master cycles from {SUPABASE_URL}/rest/v1/growout_pond_master ...")
    url_base = f"{SUPABASE_URL}/rest/v1/growout_pond_master?select=pond_index,pond,date_cycle,date_close,pond_status&order=pond,date_cycle.asc"
    all_cycles = []
    offset = 0
    limit = 1000

    while True:
        success = False
        for attempt in range(1, max_retries + 1):
            try:
                req = urllib.request.Request(f"{url_base}&limit={limit}&offset={offset}", headers=HEADERS)
                with urllib.request.urlopen(req, timeout=25) as resp:
                    batch = json.loads(resp.read().decode())
                success = True
                break
            except Exception as e:
                log(f"Cycle fetch attempt {attempt} failed: {e}. Retrying...")
                time.sleep(2 * attempt)
        if not success:
            raise RuntimeError("Failed to fetch master cycles after multiple retries.")
        if not batch:
            break
        all_cycles.extend(batch)
        offset += len(batch)
        if len(batch) < limit:
            break

    log(f"Loaded {len(all_cycles)} master cycle definitions.")

    pond_cycles = {}
    for c in all_cycles:
        p = c["pond"]
        if p not in pond_cycles:
            pond_cycles[p] = []
        pond_cycles[p].append(c)

    return pond_cycles

def resolve_pond_index(pond, log_date, pond_cycles):
    if pond not in pond_cycles:
        return None, "UNKNOWN_POND"

    cycles = pond_cycles[pond]

    # 1. Exact match within cycle bounds
    for i, c in enumerate(cycles):
        dc = c["date_cycle"]
        dclose = c["date_close"]
        if dc and dc <= log_date:
            if dclose and log_date <= dclose:
                return c["pond_index"], "EXACT"
            elif not dclose:
                next_dc = cycles[i + 1]["date_cycle"] if i + 1 < len(cycles) else None
                if not next_dc or log_date < next_dc:
                    return c["pond_index"], "OPEN_CYCLE"

    # 2. Inter-cycle preparation check
    for i in range(len(cycles) - 1):
        c1 = cycles[i]
        c2 = cycles[i + 1]
        if c1["date_close"] and c2["date_cycle"]:
            if c1["date_close"] < log_date < c2["date_cycle"]:
                return c2["pond_index"], "PRE_CYCLE_PREP"

    # 3. Fallbacks for boundary dates
    if cycles and cycles[-1]["date_cycle"] and log_date >= cycles[-1]["date_cycle"]:
        return cycles[-1]["pond_index"], "LATEST_CYCLE"
    if cycles and cycles[0]["date_cycle"] and log_date < cycles[0]["date_cycle"]:
        return cycles[0]["pond_index"], "EARLIEST_CYCLE"

    return None, "NO_CYCLE_MATCH"

def post_batch(batch, max_retries=3):
    url = f"{SUPABASE_URL}/rest/v1/lab_water_quality?on_conflict=pond_index,log_date"
    body = json.dumps(batch).encode("utf-8")
    for attempt in range(1, max_retries + 1):
        try:
            req = urllib.request.Request(url, data=body, headers=HEADERS, method="POST")
            with urllib.request.urlopen(req, timeout=30) as resp:
                if resp.status in (200, 201, 204):
                    return True
        except urllib.error.HTTPError as e:
            err_msg = e.read().decode("utf-8", errors="ignore")
            log(f"HTTP {e.code} on attempt {attempt}: {err_msg[:200]}")
            if attempt == max_retries:
                raise
        except Exception as e:
            log(f"Error on attempt {attempt}: {e}")
            if attempt == max_retries:
                raise
        time.sleep(2 * attempt)
    return False

def resolve_excel_path():
    if os.path.exists(PRIMARY_EXCEL_PATH):
        log(f"Targeting Primary Server Excel: {PRIMARY_EXCEL_PATH}")
        return PRIMARY_EXCEL_PATH
    elif os.path.exists(FALLBACK_EXCEL_PATH):
        log(f"Server drive Y: not reachable. Using fallback local copy: {FALLBACK_EXCEL_PATH}")
        return FALLBACK_EXCEL_PATH
    else:
        log(f"ERROR: Neither primary '{PRIMARY_EXCEL_PATH}' nor fallback '{FALLBACK_EXCEL_PATH}' could be found.")
        sys.exit(1)

def run_sync(is_full_sync=False, lookback_days=30):
    start_time = time.time()
    log("==================================================================")
    log("  iSHARP DAILY AUTOMATED LABORATORY WATER QUALITY SYNC")
    log(f"  Mode: {'FULL ARCHIVE (2024-2026)' if is_full_sync else f'INCREMENTAL (Recent {lookback_days} days)'}")
    log("==================================================================")

    excel_path = resolve_excel_path()
    pond_cycles = fetch_master_cycles()

    # Calculate cutoff date for incremental sync
    cutoff_date_str = None
    if not is_full_sync and lookback_days:
        cutoff_dt = datetime.date.today() - datetime.timedelta(days=lookback_days)
        cutoff_date_str = cutoff_dt.strftime("%Y-%m-%d")
        log(f"Syncing records on or after cutoff date: {cutoff_date_str}")

    # Copy to temporary file to prevent file lock contention
    temp_dir = tempfile.gettempdir()
    temp_excel = os.path.join(temp_dir, f"temp_lab_sync_{int(time.time())}.xlsx")
    log(f"Creating non-blocking temp file copy: {temp_excel} ...")
    shutil.copy2(excel_path, temp_excel)

    records_map = {}
    row_count = 0
    filtered_out_count = 0
    skipped_count = 0

    try:
        log("Opening Excel workbook (read_only=True) ...")
        wb = openpyxl.load_workbook(temp_excel, read_only=True, data_only=True)
        if "DB COMBINE" not in wb.sheetnames:
            raise ValueError(f"Sheet 'DB COMBINE' not found in workbook. Found: {wb.sheetnames}")
        sheet = wb["DB COMBINE"]

        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i == 0:
                continue  # Header
            row_count += 1

            raw_date = row[4]    # Col 4: Date
            raw_pond = row[12]   # Col 12: Pond
            if not raw_date or not raw_pond:
                skipped_count += 1
                continue

            log_date = safe_date_str(raw_date)
            pond = str(raw_pond).strip()
            if not log_date or not pond:
                skipped_count += 1
                continue

            # Apply incremental cutoff if not in full mode
            if cutoff_date_str and log_date < cutoff_date_str:
                filtered_out_count += 1
                continue

            pond_index, status = resolve_pond_index(pond, log_date, pond_cycles)
            if not pond_index:
                skipped_count += 1
                continue

            record = {
                "pond_index": pond_index,
                "pond": pond,
                "log_date": log_date,
                "doc": safe_int(row[13]),             # Col 13: DOC
                "salinity_ppt": safe_float(row[14]),  # Col 14: Salinity (ppt)
                "alkalinity": safe_float(row[18]),    # Col 18: Alkalinity
                "ammonia": safe_float(row[19]),       # Col 19: Ammonia
                "nitrite": safe_float(row[20]),       # Col 20: Nitrite
                "calcium": safe_float(row[31]),       # Col 31: Calcium
                "magnesium": safe_float(row[32]),     # Col 32: Magnesium
                "turbidity": safe_float(row[34])      # Col 34: Turbidity
            }

            dedup_key = (pond_index, log_date)
            records_map[dedup_key] = record

        wb.close()
    finally:
        if os.path.exists(temp_excel):
            try:
                os.remove(temp_excel)
            except Exception:
                pass

    total_records = list(records_map.values())
    log(f"Excel read complete. Total scanned: {row_count}, Outside lookback: {filtered_out_count}, Invalid/skipped: {skipped_count}")
    log(f"Records to upsert into Supabase: {len(total_records)}")

    if not total_records:
        log("No new/updated records to sync in the lookback window. All up to date.")
        return

    total_batches = (len(total_records) + BATCH_SIZE - 1) // BATCH_SIZE
    uploaded = 0
    for b_idx in range(total_batches):
        batch = total_records[b_idx * BATCH_SIZE : (b_idx + 1) * BATCH_SIZE]
        post_batch(batch)
        uploaded += len(batch)
        pct = (uploaded / len(total_records)) * 100
        log(f"Batch {b_idx + 1}/{total_batches} ({len(batch)} items) synced. Progress: {uploaded}/{len(total_records)} ({pct:.1f}%)")

    elapsed = time.time() - start_time
    log(f"SUCCESS: Synchronized {uploaded} records to 'lab_water_quality' in {elapsed:.1f}s.")

if __name__ == "__main__":
    is_full = "--full" in sys.argv
    run_sync(is_full_sync=is_full)
